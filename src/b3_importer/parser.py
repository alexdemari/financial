"""Parse B3 investor movement XLSX exports without trusting worksheet dimensions."""

from __future__ import annotations

import hashlib
import re
import unicodedata
import warnings
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl

from b3_importer.models import B3Movement, ParsedMovements, RejectedRow

REQUIRED_HEADERS = (
    "Entrada/Saída",
    "Data",
    "Movimentação",
    "Produto",
    "Instituição",
    "Quantidade",
    "Preço unitário",
    "Valor da Operação",
)


def parse_movements_xlsx(path: Path) -> ParsedMovements:
    """Parse the B3 ``Movimentação`` sheet into normalized movement records.

    B3 exports may declare the incorrect worksheet dimension (``A1``). Loading
    the workbook in normal mode makes openpyxl traverse the actual cells.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style.*",
            category=UserWarning,
        )
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    if "Movimentação" not in workbook.sheetnames:
        raise ValueError("B3 workbook is missing the 'Movimentação' sheet")
    worksheet = workbook["Movimentação"]
    headers = _headers(worksheet)
    missing = [header for header in REQUIRED_HEADERS if _key(header) not in headers]
    if missing:
        raise ValueError(
            f"B3 workbook is missing required columns: {', '.join(missing)}"
        )

    source_file_hash = _file_hash(path)
    imported_at = datetime.now(UTC).isoformat()
    movements: list[B3Movement] = []
    rejected_rows: list[RejectedRow] = []
    occurrence_by_fingerprint: dict[str, int] = {}
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        values = {
            header: row[column_index] if column_index < len(row) else None
            for header, column_index in headers.items()
        }
        if not any(
            value is not None and str(value).strip() for value in values.values()
        ):
            continue
        try:
            movement = _movement_from_values(
                values,
                source_file_hash=source_file_hash,
                source_row_number=row_number,
                imported_at=imported_at,
            )
        except ValueError as error:
            rejected_rows.append(RejectedRow(row_number, str(error)))
            continue
        fingerprint = movement.event_id
        occurrence = occurrence_by_fingerprint.get(fingerprint, 0) + 1
        occurrence_by_fingerprint[fingerprint] = occurrence
        movements.append(
            B3Movement(
                **{**movement.to_dict(), "event_id": f"{fingerprint}:{occurrence}"}
            )
        )
    return ParsedMovements(movements, rejected_rows)


def _headers(worksheet: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, int]:
    return {
        _key(value): index
        for index, value in enumerate(
            next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        )
        if value is not None
    }


def _movement_from_values(
    values: dict[str, object],
    *,
    source_file_hash: str,
    source_row_number: int,
    imported_at: str,
) -> B3Movement:
    event_date = _date(values[_key("Data")])
    direction = _text(values[_key("Entrada/Saída")]).lower()
    movement_type = _text(values[_key("Movimentação")])
    product = _text(values[_key("Produto")])
    institution = _text(values[_key("Instituição")])
    if not direction or not movement_type or not product or not institution:
        raise ValueError("missing required event fields")
    quantity = _number(values[_key("Quantidade")])
    unit_price_brl = _number(values[_key("Preço unitário")])
    operation_value_brl = _number(values[_key("Valor da Operação")])
    fingerprint = _fingerprint(
        event_date,
        direction,
        movement_type,
        product,
        institution,
        quantity,
        unit_price_brl,
        operation_value_brl,
    )
    return B3Movement(
        event_id=fingerprint,
        event_date=event_date,
        direction=direction,
        movement_type=movement_type,
        product=product,
        institution=institution,
        quantity=quantity,
        unit_price_brl=unit_price_brl,
        operation_value_brl=operation_value_brl,
        source_file_hash=source_file_hash,
        source_row_number=source_row_number,
        imported_at=imported_at,
    )


def _fingerprint(*values: object) -> str:
    normalized = "|".join("" if value is None else str(value) for value in values)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError("invalid event date")


def _number(value: object) -> float | None:
    if value is None or str(value).strip() in {"", "-"}:
        return None
    if isinstance(value, int | float):
        return float(value)
    text = _text(value).replace("R$", "").replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError as error:
        raise ValueError("invalid numeric value") from error


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _key(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", _text(value))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", ascii_value).strip().lower()
