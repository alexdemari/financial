"""Parser for the recorded B3 custody-position workbook schema."""

from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
import warnings
from datetime import date
from pathlib import Path

import openpyxl

from b3_importer.models import B3Position

# Recorded from B3 position export on 2026-10-12.  Keep this explicit: movement
# exports are a different schema and must never be used to infer these columns.
POSITION_SHEETS = {
    "Acoes": (
        "acoes_br",
        "Código de Negociação",
        "Preço de Fechamento",
        "Valor Atualizado",
        True,
    ),
    "BDR": (
        "bdr",
        "Código de Negociação",
        "Preço de Fechamento",
        "Valor Atualizado",
        True,
    ),
    "Empréstimos": (
        "emprestimos",
        None,
        "Preço de Fechamento",
        "Valor Atualizado",
        False,
    ),
    "Renda Fixa": (
        "renda_fixa_br",
        "Código",
        "Preço Atualizado MTM",
        "Valor Atualizado MTM",
        True,
    ),
    "Tesouro Direto": ("renda_fixa_br", "Código ISIN", None, "Valor Atualizado", True),
}


def parse_positions_xlsx(path: Path, as_of: date) -> list[B3Position]:
    """Read a B3 custody statement using its registered sheet/header contract."""
    payload = path.read_bytes()
    source_file_hash = hashlib.sha256(payload).hexdigest()
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore", message="Workbook contains no default style.*"
        )
        workbook = openpyxl.load_workbook(
            io.BytesIO(payload), read_only=False, data_only=True
        )
    try:
        missing_sheets = set(POSITION_SHEETS) - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError("B3 position workbook is missing required sheets")
        positions: list[B3Position] = []
        for sheet_name, spec in POSITION_SHEETS.items():
            positions.extend(
                _parse_sheet(workbook[sheet_name], as_of, source_file_hash, spec)
            )
        return positions
    finally:
        workbook.close()


def _parse_sheet(
    worksheet,
    as_of: date,
    source_file_hash: str,
    spec: tuple[str, str | None, str | None, str, bool],
) -> list[B3Position]:
    asset_class, code_header, _price_header, value_header, included_in_total = spec
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers: dict[str, int] = {}
    for index, value in enumerate(header_row):
        if value is None:
            continue
        if _key(value) in headers:
            raise ValueError(
                f"B3 position sheet '{worksheet.title}' has duplicate columns"
            )
        headers[_key(value)] = index
    required = ["Produto", "Instituição", "Quantidade", value_header]
    if code_header:
        required.append(code_header)
    missing = [header for header in required if _key(header) not in headers]
    if missing:
        raise ValueError(
            f"B3 position sheet '{worksheet.title}' is missing required columns"
        )
    results: list[B3Position] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        values = {
            header: row[index] if index < len(row) else None
            for header, index in headers.items()
        }
        if not any(
            value is not None and str(value).strip() for value in values.values()
        ):
            continue
        product = _text(values[_key("Produto")])
        institution = _text(values[_key("Instituição")])
        if not product or not institution:
            # B3 appends blank separators and total rows containing only a
            # value. They are not custody records and must not become a
            # synthetic position or abort the whole immutable import.
            continue
        results.append(
            B3Position(
                as_of=as_of.isoformat(),
                source_sheet=worksheet.title,
                asset_class=asset_class,
                product=product,
                institution=institution,
                trading_code=_text(values.get(_key(code_header))) or None
                if code_header
                else None,
                quantity=_number(values[_key("Quantidade")]),
                value_brl=_number(values[_key(value_header)]),
                included_in_total=included_in_total,
                source_file_hash=source_file_hash,
            )
        )
    return results


def _number(value: object) -> float:
    if value is None or str(value).strip() in {"", "-"}:
        return 0.0
    if isinstance(value, int | float):
        result = float(value)
    else:
        text = _normalize_number_text(_text(value).replace("R$", "").strip())
        try:
            result = float(text)
        except ValueError as error:
            raise ValueError("invalid numeric value in B3 position workbook") from error
    if not math.isfinite(result):
        raise ValueError("invalid numeric value in B3 position workbook")
    return result


def _normalize_number_text(text: str) -> str:
    """Convert pt-BR ('1.234,56'), thousands-only ('1.234') or plain ('10.5') text."""
    if "," in text:
        return text.replace(".", "").replace(",", ".")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):
        return text.replace(".", "")
    return text


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _key(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", _text(value))
    return (
        re.sub(r"\s+", " ", normalized.encode("ascii", "ignore").decode("ascii"))
        .strip()
        .lower()
    )
