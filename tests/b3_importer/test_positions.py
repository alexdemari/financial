from datetime import date
from pathlib import Path

import openpyxl

import pytest

from b3_importer.positions import POSITION_SHEETS, _number, parse_positions_xlsx
from b3_importer.positions_store import write_snapshot
from b3_importer.reconcile import ReconciliationRow, reconcile, write_report


def _workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (
        _asset_class,
        code,
        _price,
        value,
        _included,
    ) in POSITION_SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        headers = ["Produto", "Instituição", "Quantidade", value]
        if code:
            headers.append(code)
        sheet.append(headers)
        row = ["Produto Teste", "BTG Pactual", 10, 1000]
        if code:
            row.append("TEST3")
        sheet.append(row)
    workbook.save(path)
    return path


def test_position_parser_uses_registered_schema_and_excludes_lending(tmp_path: Path):
    positions = parse_positions_xlsx(
        _workbook(tmp_path / "posicao-2026-10-12.xlsx"), date(2026, 10, 12)
    )

    assert len(positions) == len(POSITION_SHEETS)
    assert {position.asset_class for position in positions} == {
        "acoes_br",
        "bdr",
        "emprestimos",
        "renda_fixa_br",
    }
    assert (
        next(
            position for position in positions if position.source_sheet == "Empréstimos"
        ).included_in_total
        is False
    )


def test_position_parser_ignores_b3_total_rows_without_custody_identity(tmp_path: Path):
    path = _workbook(tmp_path / "posicao-2026-10-12.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["Acoes"].append([None, None, None, 9999])
    workbook.save(path)

    positions = parse_positions_xlsx(path, date(2026, 10, 12))

    assert len(positions) == len(POSITION_SHEETS)


def test_snapshot_is_immutable_and_idempotent(tmp_path: Path):
    positions = parse_positions_xlsx(
        _workbook(tmp_path / "posicao-2026-10-12.xlsx"), date(2026, 10, 12)
    )

    first = write_snapshot(tmp_path / "snapshots", positions)
    second = write_snapshot(tmp_path / "snapshots", positions)

    assert first == second
    assert len(list((tmp_path / "snapshots").glob("*.csv"))) == 1


def test_reconciliation_reports_quantity_mismatch_without_writing_sources(
    tmp_path: Path,
):
    snapshot = tmp_path / "snapshot.csv"
    snapshot.write_text(
        "as_of,source_sheet,asset_class,product,institution,trading_code,quantity,value_brl,included_in_total,source_file_hash\n"
        "2026-10-12,Acoes,acoes_br,Produto,BTG Pactual,TEST3,10,1000,True,hash\n",
        encoding="utf-8",
    )
    btg = tmp_path / "btg.csv"
    original = "codigo,asset_type,quantidade\nTEST3,ACAO,8\n"
    btg.write_text(original, encoding="utf-8")

    rows = reconcile(snapshot, [btg])

    assert [(row.product, row.status) for row in rows] == [
        ("TEST3", "quantity_mismatch")
    ]
    assert btg.read_text(encoding="utf-8") == original


@pytest.mark.parametrize(
    ("text", "expected"),
    [("10.5", 10.5), ("1.234,56", 1234.56), ("1.234", 1234.0), ("R$ 2,5", 2.5)],
)
def test_number_parses_text_formats(text: str, expected: float):
    assert _number(text) == expected


def test_parser_rejects_duplicate_headers_and_empty_sheet(tmp_path: Path):
    path = _workbook(tmp_path / "posicao-2026-10-12.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["Acoes"]["F1"] = "Quantidade"
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate"):
        parse_positions_xlsx(path, date(2026, 10, 12))

    path = _workbook(tmp_path / "posicao-2026-10-13.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["BDR"].delete_rows(1, 2)
    workbook.save(path)
    with pytest.raises(ValueError, match="missing required columns"):
        parse_positions_xlsx(path, date(2026, 10, 13))


def test_report_names_snapshot_and_keeps_large_quantities(tmp_path: Path):
    snapshot = tmp_path / "positions_2026-10-12_abc.csv"
    report = write_report(
        [ReconciliationRow("BTG", "TEST3", 1234567, 1234568)], tmp_path, snapshot
    )
    content = report.read_text(encoding="utf-8")

    assert "Snapshot: `positions_2026-10-12_abc.csv`" in content
    assert "| 1234567 | 1234568 |" in content
